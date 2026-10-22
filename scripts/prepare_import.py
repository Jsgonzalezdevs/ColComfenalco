"""Read-only extraction. Originals and full extracted content stay outside public/.

Run with bundled Python (openpyxl, python-docx, pypdf) and xlrd 2.0.2.
Usage: python prepare_import.py SOURCE_DIRECTORY OUTPUT_DIRECTORY
"""
import base64
import datetime
import hashlib
import json
import sys
import unicodedata
import uuid
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
import xlrd
from docx import Document
from pypdf import PdfReader

NAMESPACE = uuid.UUID('f1e632de-64da-47cd-9d7e-03dc382a91af')


def identity(value):
    return str(uuid.uuid5(NAMESPACE, value))


def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value)) if unicodedata.category(c) != 'Mn').strip().lower()


def serial(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return base64.b64encode(value).decode()
    return value


def category(path):
    name = normalize(path)
    if 'actas/' in name:
        return 'Acta'
    if 'baja' in name:
        return 'Baja / solicitud'
    if 'traslado' in name:
        return 'Traslado'
    if 'no encontrados' in name:
        return 'No encontrados'
    if name.endswith('.pdf'):
        return 'Cotización'
    return 'Inventario / soporte'


def spreadsheet_rows(path):
    if path.suffix.lower() == '.xls':
        book = xlrd.open_workbook(path)
        for sheet in book.sheets():
            for index in range(sheet.nrows):
                cells = {}
                for col in range(sheet.ncols):
                    cell = sheet.cell(index, col)
                    if cell.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, book.datemode).isoformat()
                        cells[openpyxl.utils.get_column_letter(col + 1)] = {'value': serial(value)}
                if cells:
                    yield sheet.name, index + 1, cells
        return
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in formulas:
        for row_number, (row, values) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
            cells = {}
            for cell, stored in zip(row, values):
                if cell.value is not None:
                    payload = {'value': serial(stored.value if cell.data_type == 'f' else cell.value)}
                    if cell.data_type == 'f':
                        payload['formula'] = cell.value
                    cells[openpyxl.utils.get_column_letter(cell.column)] = payload
            if cells:
                yield sheet.title, row_number, cells
    formulas.close()
    cached.close()


def extract(path, ocr):
    if path.suffix.lower() in ('.xlsx', '.xls'):
        for sheet, row, cells in spreadsheet_rows(path):
            text = ' | '.join(f'{col}: {cell.get("value") if cell.get("value") is not None else cell.get("formula", "")}' for col, cell in cells.items())
            yield f'{sheet} · fila {row}', 'sheet_row', {'sheet': sheet, 'row': row, 'cells': cells}, text
    elif path.suffix.lower() == '.docx':
        doc = Document(path)
        for i, paragraph in enumerate(doc.paragraphs, 1):
            if paragraph.text.strip():
                yield f'Párrafo {i}', 'paragraph', {'paragraph': i}, paragraph.text
        for t, table in enumerate(doc.tables, 1):
            for r, row in enumerate(table.rows, 1):
                cells = [cell.text for cell in row.cells]
                if any(cell.strip() for cell in cells):
                    yield f'Tabla {t} · fila {r}', 'table_row', {'table': t, 'row': r, 'cells': cells}, ' | '.join(cells)
        # Include headers and footers, which may contain document codes and dates.
        for s, section in enumerate(doc.sections, 1):
            for label, part in [('Encabezado', section.header), ('Pie', section.footer), ('Encabezado primera página', section.first_page_header), ('Pie primera página', section.first_page_footer)]:
                texts = [p.text for p in part.paragraphs if p.text.strip()]
                texts += [' | '.join(c.text for c in row.cells) for table in part.tables for row in table.rows]
                if texts:
                    yield f'Sección {s} · {label}', 'document_metadata', {'section': s, 'part': label}, '\n'.join(texts)
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith('word/media/'):
                    sha = hashlib.sha256(archive.read(name)).hexdigest()
                    result = ocr.get(sha)
                    if result and result['lines']:
                        lines = sorted(result['lines'], key=lambda line: (-round(line['y'] / 0.012), line['x']))
                        yield f'Imagen {name}', 'ocr_image', {'image': name, 'sha256': sha, 'needs_review': True, 'lines': lines}, '\n'.join(line['text'] for line in lines)
    elif path.suffix.lower() == '.pdf':
        for i, page in enumerate(PdfReader(path).pages, 1):
            text = page.extract_text() or ''
            yield f'Página {i}', 'pdf_page', {'page': i, 'needs_ocr': not text.strip()}, text


def annotate_candidates(entries, documents):
    """Attach proposals to source rows; never resolve conflicting descriptions or custody."""
    headers = {}
    plates = {}
    doc_lookup = {d['id']: d for d in documents}
    for entry in entries:
        payload = entry['payload']
        if entry['kind'] == 'sheet_row':
            cells = {key: cell.get('value') for key, cell in payload['cells'].items()}
            group = (entry['document_id'], payload['sheet'])
        elif entry['kind'] == 'table_row':
            cells = {str(index): value for index, value in enumerate(payload['cells'])}
            group = (entry['document_id'], payload['table'])
        else:
            continue
        normalized = {key: normalize(value) for key, value in cells.items() if value is not None}
        desc_headers = {'descripcion', 'descripcion placa / activo', 'detalle del activo', 'activos', 'activo', 'activoactivos'}
        plate_headers = {'placa', 'placa inventario', 'placa / activo', 'placa/activo', 'n° placa', 'no. placa'}
        desc_keys = [key for key, value in normalized.items() if value in desc_headers]
        plate_keys = [key for key, value in normalized.items() if value in plate_headers]
        if desc_keys and (plate_keys or 'cantidad' in normalized.values()):
            headers[group] = {'descriptions': desc_keys, 'plate': plate_keys[0] if plate_keys else None,
                              'quantity': next((key for key, value in normalized.items() if value in {'cantidad', 'cant.', 'cant'}), None),
                              'status': next((key for key, value in normalized.items() if value == 'estado'), None)}
            continue
        header = headers.get(group)
        if not header:
            continue
        descriptions = [str(cells.get(key) or '').strip() for key in header['descriptions']]
        descriptions = list(dict.fromkeys(value for value in descriptions if value))
        plate = str(cells.get(header['plate']) or '').strip()
        quantity = cells.get(header['quantity'])
        if not descriptions or any(normalize(text).startswith('total') for text in descriptions):
            continue
        if not (plate.replace('.', '').isdigit() or 'sin placa' in normalize(plate) or isinstance(quantity, (int, float)) and quantity > 0):
            continue
        if len(descriptions[0]) < 3:
            continue
        flags = []
        if len({normalize(value) for value in descriptions}) > 1:
            flags.append('Las columnas de descripción no coinciden')
        if not quantity:
            flags.append('Cantidad no informada explícitamente')
        if plate.replace('.', '').isdigit() and isinstance(quantity, (int, float)) and quantity != 1:
            flags.append('La placa aparece junto a una cantidad distinta de uno')
        status = str(cells.get(header['status']) or '').strip()
        if not status:
            flags.append('Estado no informado')
        if doc_lookup[entry['document_id']]['category'] in {'Traslado', 'Baja / solicitud'}:
            flags.append('Documento de movimiento o baja; no acredita un activo adicional')
        proposal = {'plate_as_written': plate, 'descriptions': descriptions, 'quantity_as_written': quantity,
                    'status_as_written': status, 'review_flags': flags}
        payload['asset_proposal'] = proposal
        entry['kind'] = 'asset_candidate'
        if plate.replace('.', '').isdigit():
            plates.setdefault(plate, []).append(proposal)
    for proposals in plates.values():
        if len(proposals) > 1:
            for proposal in proposals:
                proposal['review_flags'].append(f'Placa repetida en {len(proposals)} filas de origen; conciliar antes de registrar')


def prepare(source, output, ocr_path=None):
    ocr = {}
    if ocr_path:
        for line in ocr_path.read_text().splitlines():
            result = json.loads(line)
            ocr[Path(result['image']).stem] = result
    documents, entries, skipped, errors = [], [], [], []
    before = {}
    for path in sorted(source.rglob('*')):
        if not path.is_file() or 'inventario-comfenalco' in path.relative_to(source).parts or output == path or output in path.parents:
            continue
        relative = path.relative_to(source).as_posix()
        if path.name.startswith('~$') or path.name == '.DS_Store':
            skipped.append(relative)
            continue
        # Archive every source file, even if its format has no text extractor.
        raw = path.read_bytes()
        checksum = hashlib.sha256(raw).hexdigest()
        before[relative] = checksum
        doc_id = identity(relative + ':' + checksum)
        doc = {'id': doc_id, 'path': relative, 'sha256': checksum, 'byte_size': len(raw), 'format': path.suffix.lower().lstrip('.'), 'category': category(relative), 'original_base64': base64.b64encode(raw).decode(), 'entry_count': 0}
        try:
            extracted = list(extract(path, ocr))
            for i, (locator, kind, payload, text) in enumerate(extracted):
                entries.append({'id': identity(doc_id + ':' + str(i)), 'document_id': doc_id, 'ordinal': i, 'locator': locator, 'kind': kind, 'payload': payload, 'text': text})
            doc['entry_count'] = len(extracted)
            doc['extraction_status'] = 'complete' if extracted else 'original_only'
        except Exception as error:
            doc['extraction_status'] = 'error'
            errors.append({'path': relative, 'error': str(error)})
        documents.append(doc)
    for relative, checksum in before.items():
        if hashlib.sha256((source / relative).read_bytes()).hexdigest() != checksum:
            raise RuntimeError('El archivo cambió durante la extracción: ' + relative)
    annotate_candidates(entries, documents)
    manifest = {'format': 'colegio-source-archive', 'version': 1, 'documents': documents, 'entries': entries}
    output.mkdir(parents=True, exist_ok=True)
    (output / 'importacion-colegio.json').write_text(json.dumps(manifest, ensure_ascii=False, allow_nan=False), encoding='utf-8')
    report = {'files': len(documents), 'by_format': dict(Counter(d['format'] for d in documents)), 'entries': len(entries), 'asset_candidate_rows': sum(e['kind']=='asset_candidate' for e in entries), 'original_bytes': sum(d['byte_size'] for d in documents), 'extraction_status': dict(Counter(d['extraction_status'] for d in documents)), 'skipped_temporary_files': skipped, 'errors': errors, 'originals_unchanged': True, 'documents': [{k: v for k, v in d.items() if k != 'original_base64'} for d in documents]}
    (output / 'verificacion-importacion.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k != 'documents'}, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit('Hay errores de extracción; los originales sí quedaron preservados.')


if __name__ == '__main__':
    prepare(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve(), Path(sys.argv[3]) if len(sys.argv)>3 else None)
